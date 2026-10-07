import os
import sys
import tempfile
import json
import openpyxl
from backend import DatabaseManager
from ai_service import AIService
from flask import Flask, jsonify, request, send_file

if getattr(sys, 'frozen', False):
    BASE_DIR = sys._MEIPASS
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

static_folder_path = os.path.join(BASE_DIR, "static")

app = Flask(__name__, static_folder=static_folder_path, static_url_path="")

db_manager = DatabaseManager()
ai_service = AIService(db_path=db_manager.session_db_path)

@app.route("/")
def index():
    return app.send_static_file("index.html")

@app.route("/api/init", methods=['POST', 'GET'])
def init_session():
    db_manager.reset_session()
    return jsonify({"status": "success", "original_filename": db_manager.original_filename})

@app.route("/api/tables", methods=["GET"])
def get_tables_list():
    try:
        staged_info = db_manager.refresh_staged_data_from_db()
        tables = list(staged_info.keys())
        return jsonify({"status": "success", "tables": tables})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/new", methods=["POST"])
def create_new():
    data = db_manager.create_empty_database()
    return jsonify({"status": "success", "data": data, "filename": "Untitled.db"})

@app.route("/api/load", methods=["POST"])
def load_db():
    if "file" not in request.files:
        return jsonify({"status": "error", "message": "No file uploaded"}), 400

    uploaded_file = request.files["file"]
    if uploaded_file.filename == "":
        return jsonify({"status": "error", "message": "No selected file"}), 400

    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".db") as temp_file:
            uploaded_file.save(temp_file.name)
            temp_path = temp_file.name

        data = db_manager.load_db_file(
            temp_path, original_filename=uploaded_file.filename
        )
        os.remove(temp_path)

        return jsonify({
            "status": "success",
            "filename": uploaded_file.filename,
            "data": data,
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/import', methods=['POST'])
def import_file():
    if 'file' not in request.files:
        return jsonify({'status': 'error', 'message': 'No file uploaded'}), 400

    file = request.files['file']
    filename = file.filename

    try:
        if filename.endswith('.xlsx') or filename.endswith('.xls'):
            wb = openpyxl.load_workbook(file, data_only=True)
            imported_tables = {}

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                all_rows = list(sheet.iter_rows(values_only=True))

                if not all_rows:
                    continue

                cols = [str(c).strip() if c is not None else "" for c in all_rows[0]]
                while cols and cols[-1] == "":
                    cols.pop()

                col_count = len(cols)
                if col_count == 0:
                    continue

                rows = []
                for row in all_rows[1:]:
                    row_data = [str(cell).strip() if cell is not None else "" for cell in row[:col_count]]
                    while len(row_data) < col_count:
                        row_data.append("")
                    rows.append(row_data)

                db_manager.update_staged_data(sheet_name, rows, cols)

                imported_tables[sheet_name] = {
                    'columns': cols,
                    'rows': rows
                }

            return jsonify({
                'status': 'success',
                'tables': imported_tables
            })

        return jsonify({'status': 'error', 'message': 'Unsupported file format'}), 400

    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)}), 500

@app.route("/api/update_staging", methods=["POST"])
def update_staging():
    payload = request.json or {}
    table_name = payload.get("table_name")
    columns = payload.get("columns", [])
    rows = payload.get("rows", [])
    page = payload.get("page", 1)
    limit = payload.get("limit")

    if not table_name:
        return jsonify(
            {"status": "error", "message": "table_name is required"}
        ), 400

    db_manager.update_staged_data(table_name, rows, columns, page=page, limit=limit)
    return jsonify({"status": "success", "message": "Staged data updated"})

@app.route("/api/save", methods=["GET", "POST"])
def save_database():
    try:
        temp_dir = tempfile.mkdtemp()
        filename = db_manager.original_filename or "database.db"
        save_path = os.path.join(temp_dir, filename)

        db_manager.save_to_db(save_path)

        return send_file(
            save_path,
            as_attachment=True,
            download_name=filename,
            mimetype="application/x-sqlite3",
        )

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/delete_table", methods=["POST"])
def delete_table():
    try:
        data = request.get_json() or {}
        table_name = data.get("table_name")

        if not table_name:
            return jsonify(
                {"status": "error", "message": "table_name is required"}
            ), 400

        success = db_manager.delete_table_from_staging(table_name)
        if success:
            return jsonify({"status": "success"})
        return jsonify(
            {"status": "error", "message": f"Table '{table_name}' not found"}
        ), 404
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/export_xlsx", methods=["GET", "POST"])
def export_xlsx():
    try:
        saved_path = db_manager.export_xlsx()

        return jsonify({
            "status": "success",
            "message": f"Exported successfully to Downloads!",
            "path": saved_path
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/schema", methods=["GET"])
def get_schema():
    try:
        summary = db_manager.get_schema_summary()
        return jsonify({"status": "success", "schema": summary})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/rename_table", methods=["POST"])
def rename_table():
    try:
        data = request.get_json() or {}
        old_name = data.get("old_name")
        new_name = data.get("new_name")

        if not old_name or not new_name:
            return jsonify(
                {"status": "error", "message": "old_name and new_name are required"}
            ), 400

        success = db_manager.rename_table_in_staging(old_name, new_name)
        if success:
            return jsonify({"status": "success"})
        return jsonify(
            {"status": "error", "message": f"Table '{old_name}' could not be renamed"}
        ), 400
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/logs", methods=["GET"])
def get_logs():
    try:
        limit = request.args.get("limit", default=50, type=int)
        logs = db_manager.get_recent_logs(limit=limit)
        return jsonify({"status": "success", "logs": logs})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/data", methods=["GET"])
def get_table_data():
    table_name = request.args.get('table')
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 100))
    
    sort_orders_raw = request.args.get('sort_orders')
    sort_orders = None
    if sort_orders_raw:
        try:
            sort_orders = json.loads(sort_orders_raw)
        except json.JSONDecodeError:
            pass

    if not table_name:
        return jsonify({"status": "error", "message": "Table name required"}), 400

    data = db_manager.get_table_page(table_name, page, limit, sort_orders=sort_orders)
    data["status"] = "success"
    return jsonify(data)
    
@app.route("/api/search", methods=["POST"])
def search():
    data = request.get_json() or {}
    query = data.get("query", "").strip()
    match_case = data.get("match_case", False)
    match_word = data.get("match_word", False)
    scope = data.get("scope", {})  

    if not query or not scope:
        return jsonify({"status": "success", "matches": []})

    matches = db_manager.search_tables(query, scope, match_case, match_word)
    return jsonify({"status": "success", "matches": matches})

@app.route("/api/ai/query", methods=["POST"])
def ai_query():
    try:
        data = request.get_json() or {}
        query_text = data.get("query", "")
        active_table = data.get("active_table", "Table1")

        if not query_text:
            return jsonify({"status": "error", "message": "Query text is required"}), 400

        response = ai_service.process_query(query_text=query_text, active_table=active_table)
        response['status'] = 'success'
        return jsonify(response)

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/ai/confirm", methods=["POST"])
def ai_confirm():
    try:
        data = request.get_json() or {}
        action_payload = data.get("payload")
        accepted = data.get("accepted", False)

        if not action_payload:
            return jsonify({"status": "error", "message": "Action payload is required"}), 400

        if not accepted:
            return jsonify({
                "status": "success",
                "message": "Action canceled by user. No changes were applied.",
                "actions": []
            })

        response = ai_service.confirm_action(action_payload)
        db_manager.refresh_staged_data_from_db()

        response['status'] = 'success'
        return jsonify(response)

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route("/api/delete_rows", methods=["POST"])
def delete_rows():
    try:
        data = request.get_json() or {}
        table_name = data.get("table_name")
        row_indices = data.get("rows", [])
        page = data.get("page", 1)
        limit = data.get("limit", 100)

        if not table_name or not row_indices:
            return jsonify({"status": "error", "message": "table_name and rows are required"}), 400

        success = db_manager.delete_rows_from_staging(
            table_name=table_name,
            row_indices=row_indices,
            page=page,
            limit=limit
        )
        if success:
            return jsonify({"status": "success"})
        return jsonify({"status": "error", "message": f"Table '{table_name}' not found"}), 404
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

if __name__ == "__main__":
    app.run(debug=True, port=5000)