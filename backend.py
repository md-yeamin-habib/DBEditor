import os
import re
from pathlib import Path
import sqlite3
import openpyxl
import shutil 
import logging
import sys

# Custom Handler that forces Python to write log lines to disk instantly
class FlushingFileHandler(logging.FileHandler):
    def emit(self, record):
        super().emit(record)
        self.flush()

# Clean HH:MM:SS time format without milliseconds
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
    handlers=[
        FlushingFileHandler("sql_queries.log", mode="a", encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
def log_sql(query):
    """Callback function to capture and format executed SQL statements (ignoring SELECT & PRAGMA)."""
    clean_query = " ".join(query.split())
    
    upper_query = clean_query.upper()
    if upper_query.startswith("SELECT") or upper_query.startswith("PRAGMA"):
        return

    logging.info(f"EXEC SQL: {clean_query}")

class DatabaseManager:

    def __init__(self, work_dir="session_cache"):
        self.work_dir = work_dir
        os.makedirs(self.work_dir, exist_ok=True)

        self.session_db_path = os.path.join(self.work_dir, "working_session.db")
        self.original_filename = "database.db"
        self.staged_data = {}

    def _get_connection(self):
        """Returns a SQLite connection with trace logging attached."""
        conn = sqlite3.connect(self.session_db_path)
        conn.set_trace_callback(log_sql)
        return conn

    def reset_session(self):
        """Removes the working session database and clears staged state."""
        if os.path.exists(self.session_db_path):
            try:
                os.remove(self.session_db_path)
            except Exception as e:
                logging.error(f"Failed to delete session DB: {e}")
        self.original_filename = "new_database.db"
        self.staged_data = {}
        return self.staged_data

    def rename_table_in_staging(self, old_name: str, new_name: str) -> bool:
        """Renames a table in SQLite directly and updates staged_data."""
        if not old_name or not new_name or old_name == new_name:
            return False

        conn = self._get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute(f'ALTER TABLE "{old_name}" RENAME TO "{new_name}";')
            conn.commit()
            if old_name in self.staged_data:
                self.staged_data[new_name] = self.staged_data.pop(old_name)
            return True
        except Exception as e:
            conn.rollback()
            raise e
        finally:
            conn.close()

    def delete_table_from_staging(self, table_name: str) -> bool:
        """Drops table from session SQLite database with SQL trace logging attached."""
        conn = self._get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute(f'DROP TABLE IF EXISTS "{table_name}";')
            conn.commit()
            if table_name in self.staged_data:
                del self.staged_data[table_name]
            return True
        except Exception as e:
            conn.rollback()
            raise e
        finally:
            conn.close()

    def update_staged_data(self, table_name: str, rows: list, columns: list, page: int = 1, limit: int = None):
        """
        Executes targeted SQL operations against the active working session SQLite database file.
        Detects blank/empty rows and explicitly deletes their corresponding _rowid_ entries via DELETE WHERE _rowid_ IN (...).
        """
        conn = self._get_connection()
        cursor = conn.cursor()

        try:
            col_count = len(columns)

            cursor.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name=?;",
                (table_name,),
            )
            table_exists = cursor.fetchone() is not None

            if not table_exists:
                # 1. CREATE TABLE
                col_defs = []
                for c_idx, col in enumerate(columns):
                    col_vals = [r[c_idx] for r in rows if len(r) > c_idx]
                    col_defs.append(f'"{col}" {self.infer_column_type(col_vals)}')

                cursor.execute(
                    f'CREATE TABLE "{table_name}" (_rowid_ INTEGER PRIMARY KEY AUTOINCREMENT, {", ".join(col_defs)});'
                )

                # Insert non-empty rows only
                valid_rows = []
                for r in rows:
                    sanitized_row = [str(x).strip() if x is not None else "" for x in r[:col_count]]
                    while len(sanitized_row) < col_count:
                        sanitized_row.append("")
                    if any(cell != "" for cell in sanitized_row):
                        valid_rows.append(sanitized_row)

                if valid_rows:
                    col_names = ", ".join([f'"{c}"' for c in columns])
                    placeholders = ", ".join(["?"] * col_count)
                    cursor.executemany(
                        f'INSERT INTO "{table_name}" ({col_names}) VALUES ({placeholders});',
                        valid_rows,
                    )
            else:
                existing_cols = self._get_user_columns(cursor, table_name)

                # 2. ADD NEW COLUMNS
                for c_idx, col in enumerate(columns):
                    if col not in existing_cols:
                        col_vals = [r[c_idx] for r in rows if len(r) > c_idx]
                        cursor.execute(
                            f'ALTER TABLE "{table_name}" ADD COLUMN "{col}" {self.infer_column_type(col_vals)};'
                        )

                # 3. DROP DELETED COLUMNS
                for col in existing_cols:
                    if col not in columns:
                        cursor.execute(f'ALTER TABLE "{table_name}" DROP COLUMN "{col}";')

                col_names_str = ", ".join([f'"{c}"' for c in columns]) if columns else ""

                if col_names_str:
                    # Fetch existing database rows for this scope
                    query = f'SELECT _rowid_, {col_names_str} FROM "{table_name}" ORDER BY _rowid_ ASC'
                    if limit is not None:
                        offset = (page - 1) * limit
                        query += f' LIMIT {int(limit)} OFFSET {int(offset)}'
                
                    cursor.execute(query)
                    db_data = cursor.fetchall()
                    db_rowid_s = [r[0] for r in db_data]
                    db_rows = [list(r[1:]) for r in db_data]

                    rows_to_delete = []
                    
                    # 4. PROCESS ROWS FOR UPDATE, INSERT, OR DELETE
                    for idx, r in enumerate(rows):
                        sanitized_row = [str(x).strip() if x is not None else "" for x in r[:col_count]]
                        while len(sanitized_row) < col_count:
                            sanitized_row.append("")

                        is_blank = all(cell == "" for cell in sanitized_row)

                        if idx < len(db_rowid_s):
                            row_id = db_rowid_s[idx]
                            
                            if is_blank:
                                # Row cleared in UI -> mark _rowid_ for deletion
                                rows_to_delete.append(row_id)
                            else:
                                old_row = [str(x).strip() if x is not None else "" for x in db_rows[idx]]
                                if old_row != sanitized_row:
                                    changed_cols = []
                                    changed_vals = []
                                    for c_idx, col in enumerate(columns):
                                        old_val = old_row[c_idx] if c_idx < len(old_row) else ""
                                        new_val = sanitized_row[c_idx]
                                        if old_val != new_val:
                                            changed_cols.append(f'"{col}" = ?')
                                            changed_vals.append(new_val)

                                    if changed_cols:
                                        set_clause = ", ".join(changed_cols)
                                        cursor.execute(
                                            f'UPDATE "{table_name}" SET {set_clause} WHERE _rowid_ = ?;',
                                            (*changed_vals, row_id),
                                        )
                        else:
                            # New row added at the bottom (ignore if blank)
                            if not is_blank:
                                placeholders = ", ".join(["?"] * col_count)
                                cursor.execute(
                                    f'INSERT INTO "{table_name}" ({col_names_str}) VALUES ({placeholders});',
                                    sanitized_row,
                                )

                    # Check for excess rows removed from page trailing end
                    if len(rows) < len(db_rowid_s):
                        excess_ids = db_rowid_s[len(rows):]
                        rows_to_delete.extend(excess_ids)

                    # 5. EXECUTE BULK DELETE FOR ALL IDENTIFIED BLANK OR REMOVED ROWIDs
                    if rows_to_delete:
                        placeholders = ", ".join(["?"] * len(rows_to_delete))
                        cursor.execute(
                            f'DELETE FROM "{table_name}" WHERE _rowid_ IN ({placeholders});',
                            rows_to_delete,
                        )

            conn.commit()
        except Exception as e:
            conn.rollback()
            raise e
        finally:
            conn.close()

        self.refresh_staged_data_from_db()
        return True

    def get_recent_logs(self, limit: int = 50) -> list:
        """Reads and returns only pure SQL query logs from sql_queries.log."""
        log_file = "sql_queries.log"
        if not os.path.exists(log_file):
            return []

        try:
            with open(log_file, "r", encoding="utf-8", errors="ignore") as f:
                lines = f.readlines()

            filtered_logs = []
            for line in lines:
                if "EXEC SQL:" in line:
                    sql_part = line.split("EXEC SQL:", 1)[1].strip().upper()
                    if not (sql_part.startswith("SELECT") or sql_part.startswith("PRAGMA")):
                        filtered_logs.append(line.strip())

            return filtered_logs[-limit:]
        except Exception as e:
            logging.error(f"Error reading logs: {e}")
            return []

    def get_downloads_path(self) -> Path:
        return Path.home() / "Downloads"

    def create_empty_database(self):
        if os.path.exists(self.session_db_path):
            os.remove(self.session_db_path)
        self.original_filename = "new_database.db"
        self.staged_data = {}
        return self.staged_data

    def _get_user_columns(self, cursor, table_name: str) -> list:
        cursor.execute(f'PRAGMA table_info("{table_name}");')
        cols_info = cursor.fetchall()
    
        user_cols = []
        for col in cols_info:
            col_name = col[1]
            is_pk = col[5] == 1
            col_type = col[2].upper()
        
            if col_name.lower() in ("_row_id", "_rowid_", "rowid", "oid"):
                continue
            
            if is_pk and col_type in ("INTEGER", "INT", "BIGINT", "SMALLINT"):
                continue
            
            user_cols.append(col_name)
        
        return user_cols

    @staticmethod
    def infer_column_type(values: list) -> str:
        non_empty = [str(v).strip() for v in values if v is not None and str(v).strip() != ""]
        if not non_empty:
            return "TEXT"

        is_int = True
        is_float = True

        for val in non_empty:
            if not (val.isdigit() or (val.startswith("-") and val[1:].isdigit())):
                is_int = False
            try:
                float(val)
            except ValueError:
                is_float = False

        if is_int:
            return "INTEGER"
        if is_float:
            return "REAL"
        return "TEXT"

    def load_db_file(self, uploaded_file_path: str, original_filename: str = None):
        if not os.path.exists(uploaded_file_path):
            raise FileNotFoundError(f"Database file not found: {uploaded_file_path}")

        self.original_filename = original_filename or os.path.basename(uploaded_file_path)

        if os.path.exists(self.session_db_path):
            os.remove(self.session_db_path)

        with open(uploaded_file_path, "rb") as src, open(self.session_db_path, "wb") as dst:
            dst.write(src.read())

        return self.refresh_staged_data_from_db()

    def refresh_staged_data_from_db(self):
        self.staged_data = {}
        conn = self._get_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';")
        tables = [row[0] for row in cursor.fetchall()]

        for table in tables:
            columns = self._get_user_columns(cursor, table)
            cursor.execute(f'SELECT COUNT(*) FROM "{table}";')
            total_rows = cursor.fetchone()[0]
            self.staged_data[table] = {"columns": columns, "total_rows": total_rows}

        conn.close()
        return self.staged_data

    def save_to_db(self, save_path: str = None) -> str:
        if not save_path or not os.path.isabs(save_path):
            filename = save_path if save_path else self.original_filename
            save_path = str(self.get_downloads_path() / filename)

        os.makedirs(os.path.dirname(os.path.abspath(save_path)), exist_ok=True)

        if os.path.exists(self.session_db_path):
            shutil.copyfile(self.session_db_path, save_path)

        return save_path

    def export_xlsx(self, export_path: str = None) -> str:
        if not export_path or not os.path.isabs(export_path):
            filename = (
                self.original_filename.replace(".db", ".xlsx")
                if self.original_filename
                else "exported_database.xlsx"
            )
            export_path = str(self.get_downloads_path() / filename)

        os.makedirs(os.path.dirname(os.path.abspath(export_path)), exist_ok=True)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        conn = self._get_connection()
        cursor = conn.cursor()

        try:
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';")
            tables = [row[0] for row in cursor.fetchall()]

            if not tables:
                wb.create_sheet(title="Sheet1")
            else:
                for table_name in tables:
                    ws = wb.create_sheet(title=table_name)

                    cursor.execute(f'PRAGMA table_info("{table_name}");')
                    columns = self._get_user_columns(cursor, table_name)

                    if columns:
                        ws.append(columns)

                        col_select = ", ".join([f'"{c}"' for c in columns])
                        cursor.execute(f'SELECT {col_select} FROM "{table_name}" ORDER BY _rowid_ ASC;')
                        all_rows = cursor.fetchall()

                        for row in all_rows:
                            ws.append(list(row))

            wb.save(export_path)
            return export_path
        finally:
            conn.close()

    def import_xlsx_to_table(self, xlsx_path: str, target_table: str) -> dict:
        """Parses Excel data and merges or creates a staged table."""
        if not os.path.exists(xlsx_path):
            raise FileNotFoundError(f"Excel file not found: {xlsx_path}")

        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        sheet = wb.active

        all_rows = list(sheet.iter_rows(values_only=True))
        if not all_rows:
            raise ValueError("The imported Excel file is empty.")

        imported_columns = [
            str(c).strip() if c is not None else "" for c in all_rows[0]
        ]
        while imported_columns and imported_columns[-1] == "":
            imported_columns.pop()

        col_count = len(imported_columns)
        if col_count == 0:
            raise ValueError("No valid column headers found in Excel file.")

        imported_data_rows = []
        for row in all_rows[1:]:
            row_data = [
                str(val) if val is not None else "" for val in row[:col_count]
            ]
            while len(row_data) < col_count:
                row_data.append("")

            if any(cell.strip() != "" for cell in row_data):
                imported_data_rows.append(row_data)

        if target_table not in self.staged_data:
            self.staged_data[target_table] = {"columns": [], "rows": []}

        current_table = self.staged_data[target_table]
        existing_cols = [
            c.strip() for c in current_table.get("columns", []) if c.strip()
        ]

        if not existing_cols:
            self.staged_data[target_table] = {
                "columns": imported_columns,
                "rows": imported_data_rows,
            }
        else:
            norm_existing = [c.lower() for c in existing_cols]
            norm_imported = [c.lower() for c in imported_columns]

            if norm_existing != norm_imported:
                raise ValueError(
                    f"Header mismatch! Active table '{target_table}' has columns: {existing_cols}, "
                    f"but imported Excel file has: {imported_columns}"
                )

        self.update_staged_data(
            target_table,
            imported_data_rows,
            imported_columns if not existing_cols else existing_cols,
        )

        return self.staged_data[target_table]
    
    def get_schema_summary(self) -> dict:
        """Provides table structures, inferred types, and sample data for AI context."""
        summary = {}
        for table_name, content in self.staged_data.items():
            cols = content.get("columns", [])
            rows = content.get("rows", [])
            summary[table_name] = {
                "columns": cols,
                "column_types": [
                    self.infer_column_type([r[i] for r in rows if len(r) > i])
                    for i in range(len(cols))
                ],
                "row_count": len(rows),
                "sample_data": rows[:3],
            }
        return summary

    def get_table_page(self, table_name: str, page: int = 1, limit: int = 100, sort_orders: dict = None) -> dict:
        conn = self._get_connection()
        cursor = conn.cursor()

        try:
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name=?;", (table_name,))
            if not cursor.fetchone():
                return {"columns": [], "rows": [], "total_rows": 0, "page": page, "total_pages": 0}

            cursor.execute(f'SELECT COUNT(*) FROM "{table_name}"')
            total_rows = cursor.fetchone()[0]

            columns = self._get_user_columns(cursor, table_name)

            offset = (page - 1) * limit
            rows = []
            if columns:
                col_select = ", ".join([f'"{c}"' for c in columns])
            
                order_parts = []
                if sort_orders and isinstance(sort_orders, dict):
                    for col_name, direction in sort_orders.items():
                        if col_name in columns:
                            sort_dir = "DESC" if str(direction).upper() == "DESC" else "ASC"
                            order_parts.append(f'"{col_name}" {sort_dir}')

                if order_parts:
                    order_parts.append('_rowid_ ASC')
                    order_clause = f"ORDER BY {', '.join(order_parts)}"
                else:
                    order_clause = 'ORDER BY _rowid_ ASC'

                cursor.execute(
                    f'SELECT {col_select} FROM "{table_name}" {order_clause} LIMIT ? OFFSET ?;',
                    (limit, offset)
                )
                rows = [list(row) for row in cursor.fetchall()]

            total_pages = (total_rows + limit - 1) // limit if limit > 0 else 1

            return {
                "columns": columns,
                "rows": rows,
                "total_rows": total_rows,
                "page": page,
                "limit": limit,
                "total_pages": max(1, total_pages)
            }
        finally:
            conn.close()

    def search_tables(self, query_str: str, scope: dict, match_case: bool = False, match_word: bool = False):
        matches = []
        page_size = 100
        conn = self._get_connection()
        cursor = conn.cursor()

        if not match_case:
            conn.execute("PRAGMA case_sensitive_like = OFF;")
        else:
            conn.execute("PRAGMA case_sensitive_like = ON;")

        subqueries = []
        params = []

        try:
            for table_name, columns in scope.items():
                if not columns:
                    continue

                for col in columns:
                    if match_word:
                        subqueries.append(f'''
                            SELECT '{table_name}' AS table_name, 
                                   db_rowid, 
                                   visual_pos, 
                                   '{col}' AS col_name, 
                                   cell_val 
                            FROM (
                                SELECT _rowid_ AS db_rowid,
                                       "{col}" AS cell_val,
                                       ROW_NUMBER() OVER (ORDER BY _rowid_ ASC) AS visual_pos
                                FROM "{table_name}"
                            ) 
                            WHERE cell_val REGEXP ?
                        ''')
                        params.append(rf'\b{re.escape(query_str)}\b')
                    else:
                        subqueries.append(f'''
                            SELECT '{table_name}' AS table_name, 
                                   db_rowid, 
                                   visual_pos, 
                                   '{col}' AS col_name, 
                                   cell_val 
                            FROM (
                                SELECT _rowid_ AS db_rowid,
                                       "{col}" AS cell_val,
                                       ROW_NUMBER() OVER (ORDER BY _rowid_ ASC) AS visual_pos
                                FROM "{table_name}"
                            ) 
                            WHERE cell_val LIKE ?
                        ''')
                        params.append(f"%{query_str}%")

            if not subqueries:
                return []

            full_sql = " UNION ALL ".join(subqueries) + " ORDER BY db_rowid ASC;"
            cursor.execute(full_sql, params)
            results = cursor.fetchall()

            for table_name, db_rowid, visual_pos, col_name, cell_val in results:
                col_index = scope[table_name].index(col_name) if table_name in scope else 0
                
                calc_page = ((visual_pos - 1) // page_size) + 1
                row_in_page = ((visual_pos - 1) % page_size) + 1

                matches.append({
                    'table': table_name,
                    'row': row_in_page,
                    'col_name': col_name,
                    'col': col_index,
                    'page': calc_page
                })

        except sqlite3.Error as e:
            logging.error(f"SQL Search Error: {e}")
        finally:
            conn.close()

        return matches

    def delete_rows_from_staging(self, table_name: str, row_indices: list, page: int = 1, limit: int = 100) -> bool:
        """
        Deletes specific rows from a table using their relative visual positions on a given page.
        """
        conn = self._get_connection()
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name=?;", (table_name,))
            if not cursor.fetchone():
                return False

            columns = self._get_user_columns(cursor, table_name)
            if not columns:
                return False

            col_names_str = ", ".join([f'"{c}"' for c in columns])
            offset = (page - 1) * limit

            # Retrieve database _rowid_s for the current page
            query = f'SELECT _rowid_ FROM "{table_name}" ORDER BY _rowid_ ASC LIMIT ? OFFSET ?;'
            cursor.execute(query, (limit, offset))
            db_data = cursor.fetchall()
            db_rowids = [r[0] for r in db_data]

            # Map 1-based relative row index on current page to database _rowid_
            rowids_to_delete = []
            for r_idx in row_indices:
                # Convert 1-based row relative to page offset
                relative_idx = r_idx - 1 - offset if r_idx > offset else r_idx - 1
                if 0 <= relative_idx < len(db_rowids):
                    rowids_to_delete.append(db_rowids[relative_idx])

            if rowids_to_delete:
                placeholders = ", ".join(["?"] * len(rowids_to_delete))
                cursor.execute(
                    f'DELETE FROM "{table_name}" WHERE _rowid_ IN ({placeholders});',
                    rowids_to_delete,
                )
                conn.commit()

            self.refresh_staged_data_from_db()
            return True
        except Exception as e:
            conn.rollback()
            raise e
        finally:
            conn.close()
